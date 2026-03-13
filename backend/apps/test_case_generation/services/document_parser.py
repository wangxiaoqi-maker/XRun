"""文档解析服务 - 支持 PDF / DOCX / TXT / MD，集成图片分析"""
import os
import tempfile
import shutil
from dataclasses import dataclass, field
from typing import Optional
from pathlib import Path

from ..core.logger import get_logger

logger = get_logger("doc_parser_svc")


@dataclass
class ParsedDocument:
    """解析结果，包含文本和图片分析"""
    text: str
    image_descriptions: list[dict] = field(default_factory=list)
    image_error: str | None = None


class DocumentParserService:

    async def parse_file(self, file_path: str, vision_model_id: str = None, llm_service=None) -> ParsedDocument:
        """
        解析文档（async），支持图片 LLM 分析。
        vision_model_id + llm_service 同时存在时才调用视觉模型分析图片。
        """
        ext = os.path.splitext(file_path)[1].lower()
        parsers = {
            ".pdf": self._parse_pdf,
            ".docx": self._parse_docx,
            ".doc": self._parse_docx,
            ".txt": self._parse_text,
            ".md": self._parse_text,
        }
        parser = parsers.get(ext)
        if not parser:
            raise ValueError(f"不支持的文件类型: {ext}")

        text = parser(file_path)

        image_descriptions = []
        image_error = None
        if vision_model_id and llm_service:
            try:
                image_descriptions = await self._extract_and_analyze_images(
                    file_path, vision_model_id, llm_service
                )
                if image_descriptions:
                    logger.info(f"[DocumentParser] {file_path}: 分析了 {len(image_descriptions)} 张图片")
                else:
                    logger.warning(f"[DocumentParser] {file_path}: 未提取到有效图片描述")
            except Exception as e:
                image_error = str(e)
                logger.error(f"[DocumentParser] 图片分析异常: {e}")

        result = ParsedDocument(text=text, image_descriptions=image_descriptions)
        result.image_error = image_error
        return result

    # --- Private parsers (同步，纯文本提取) ---

    def _parse_pdf(self, file_path: str) -> str:
        from PyPDF2 import PdfReader

        reader = PdfReader(file_path)
        pages = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[第{i}页]\n{text.strip()}")
        return "\n\n".join(pages)

    def _parse_docx(self, file_path: str) -> str:
        from docx import Document

        doc = Document(file_path)
        sections = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            if para.style and para.style.name.startswith("Heading"):
                sections.append(f"\n## {text}")
            else:
                sections.append(text)

        tables_text = self._extract_docx_tables(doc)
        if tables_text:
            sections.append("\n## 表格内容\n" + tables_text)

        return "\n".join(sections)

    def _extract_docx_tables(self, doc) -> str:
        tables = []
        for i, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                header = rows[0]
                sep = " | ".join(["---"] * len(table.rows[0].cells))
                body = "\n".join(rows[1:])
                tables.append(f"{header}\n{sep}\n{body}")
        return "\n\n".join(tables)

    def _parse_text(self, file_path: str) -> str:
        encodings = ["utf-8", "gbk", "gb2312", "latin-1"]
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        raise ValueError(f"无法解码文件: {file_path}")

    # --- 图片提取和分析（async，真正调用 LLM） ---

    async def _extract_and_analyze_images(
        self, file_path: str, vision_model_id: str, llm_service
    ) -> list[dict]:
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            return await self._extract_pdf_images(file_path, vision_model_id, llm_service)
        elif ext in (".docx", ".doc"):
            return await self._extract_docx_images(file_path, vision_model_id, llm_service)
        else:
            return []

    async def _extract_pdf_images(
        self, file_path: str, vision_model_id: str, llm_service
    ) -> list[dict]:
        from PyPDF2 import PdfReader
        from PIL import Image
        import io

        descriptions = []
        temp_dir = tempfile.mkdtemp()

        try:
            reader = PdfReader(file_path)
            for page_num, page in enumerate(reader.pages, 1):
                if "/XObject" not in (page.get("/Resources") or {}):
                    continue

                xobjects = page["/Resources"]["/XObject"].get_object()
                for obj_name in xobjects:
                    obj = xobjects[obj_name]
                    if obj.get("/Subtype") == "/Image":
                        try:
                            data = obj.get_data()
                            img = Image.open(io.BytesIO(data))

                            if img.width < 50 or img.height < 50:
                                continue

                            img_path = os.path.join(temp_dir, f"page{page_num}_{obj_name[1:]}.png")
                            img.save(img_path)

                            logger.info(f"[DocumentParser] PDF 图片 p{page_num}/{obj_name}: {img.width}x{img.height}")

                            description = await self._analyze_image(
                                img_path, vision_model_id, llm_service
                            )
                            if description:
                                descriptions.append({
                                    "page": page_num,
                                    "object_name": obj_name,
                                    "description": description,
                                })
                        except Exception as e:
                            err_str = str(e)
                            fatal_keywords = ("403", "401", "429", "InvalidToken", "Unauthorized",
                                              "suspended", "insufficient", "quota", "exceeded")
                            if any(kw in err_str for kw in fatal_keywords):
                                raise RuntimeError(f"视觉模型调用失败（API 错误），请检查模型配置和余额: {err_str}") from e
                            logger.warning(f"[DocumentParser] PDF 图片提取失败 p{page_num}/{obj_name}: {e}")
                            continue
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

        return descriptions

    async def _extract_docx_images(
        self, file_path: str, vision_model_id: str, llm_service
    ) -> list[dict]:
        from docx import Document

        descriptions = []
        temp_dir = tempfile.mkdtemp()

        try:
            doc = Document(file_path)
            image_index = 0

            for rel in doc.part.rels.values():
                if "image" in rel.reltype:
                    try:
                        image_part = rel.target_part
                        image_data = image_part.blob

                        content_type = getattr(image_part, 'content_type', 'image/png')
                        ext = ".png"
                        if "jpeg" in content_type or "jpg" in content_type:
                            ext = ".jpg"

                        img_path = os.path.join(temp_dir, f"image_{image_index}{ext}")
                        with open(img_path, "wb") as f:
                            f.write(image_data)

                        from PIL import Image
                        img = Image.open(img_path)
                        if img.width < 50 or img.height < 50:
                            image_index += 1
                            continue

                        logger.info(f"[DocumentParser] 分析图片 {image_index}: {img.width}x{img.height}")

                        description = await self._analyze_image(
                            img_path, vision_model_id, llm_service
                        )
                        if description:
                            descriptions.append({
                                "page": 0,
                                "index": image_index,
                                "description": description,
                            })

                        image_index += 1
                    except Exception as e:
                        err_str = str(e)
                        fatal_keywords = ("403", "401", "429", "InvalidToken", "Unauthorized",
                                          "suspended", "insufficient", "quota", "exceeded")
                        if any(kw in err_str for kw in fatal_keywords):
                            raise RuntimeError(f"视觉模型调用失败（API 错误），请检查模型配置和余额: {err_str}") from e
                        logger.warning(f"[DocumentParser] 提取图片 {image_index} 失败: {e}")
                        image_index += 1
                        continue
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

        return descriptions

    async def _analyze_image(
        self, image_path: str, vision_model_id: str, llm_service
    ) -> Optional[str]:
        """用视觉模型分析单张图片（async，真正调用 LLM）"""
        system_prompt = """你是文档分析专家。请详细描述图片中的内容，包括：
1. 图表类型（流程图、架构图、截图、表格等）
2. 关键元素和标注
3. 与测试相关的信息（如果有）

请用简洁的中文描述。"""

        user_prompt = "请分析这张图片，描述其内容和含义。"

        result = await llm_service.chat_with_vision(
            image_path=image_path,
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            model_id=vision_model_id,
        )

        return result if result else None
